#!/usr/bin/env python3
"""Sprint 125 R1: synthetic PowerPoint corpus (python-pptx 1.0.2, 16:9).

    python3 gen-fixtures.py

Writes fixtures/NN-*.pptx (one area each) and the failure fixtures f1-f4 and f6 (f6 is a
Word file, built from Sprint 124's corpus). f5 is written by PowerPoint: ppt-password.sh.
Every archive is rewritten with fixed timestamps so reruns are byte-identical.
Nothing here comes from a real document. A research script: nothing ships from it.
"""
import copy
import datetime
import io
import math
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / 'fixtures'
OUT.mkdir(exist_ok=True)
W, H = 12192000, 6858000  # 13.333 in x 7.5 in = 1280 x 720 px at 96 dpi
FIXED = datetime.datetime(2026, 1, 1, 0, 0, 0)

INDIGO = RGBColor(0x3B, 0x3F, 0xB5)
TEAL = RGBColor(0x0F, 0x9D, 0x8A)
AMBER = RGBColor(0xE8, 0xA3, 0x17)
ROSE = RGBColor(0xD1, 0x3B, 0x5C)
GREY = RGBColor(0x55, 0x5B, 0x66)


def new_deck(title):
    prs = Presentation()
    k = W / prs.slide_width
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    # The default template is 4:3; stretch the master's and layouts' placeholders to 16:9.
    # Only shapes with their own xfrm: a layout placeholder without one inherits the master's
    # (already stretched) position, and python-pptx's setters would write a partial xfrm at y=0.
    for owner in [prs.slide_master, *prs.slide_layouts]:
        for sh in owner.shapes:
            for xfrm in sh._element.iter(qn('a:xfrm')):
                off, ext = xfrm.find(qn('a:off')), xfrm.find(qn('a:ext'))
                if off is not None and ext is not None:
                    off.set('x', str(int(int(off.get('x')) * k)))
                    ext.set('cx', str(int(int(ext.get('cx')) * k)))
                break
    cp = prs.core_properties
    cp.title, cp.author, cp.last_modified_by = title, 'Ritemark corpus', 'Ritemark corpus'
    cp.created = cp.modified = FIXED
    cp.revision = 1
    return prs


def save(prs, name):
    buf = io.BytesIO()
    prs.save(buf)
    normalise(buf.getvalue(), OUT / name)


def rezip(blob):
    """Rewrite a ZIP with fixed entry timestamps; embedded chart workbooks (.xlsx) recursively,
    with their core-properties dates pinned (XlsxWriter stamps the current time)."""
    import re
    src = zipfile.ZipFile(io.BytesIO(blob))
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename.endswith('.xlsx'):
                data = rezip(data)
            elif info.filename == 'docProps/core.xml':
                data = re.sub(rb'(<dcterms:(created|modified)[^>]*>)[^<]*', rb'\g<1>2026-01-01T00:00:00Z', data)
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(zi, data)
    return out.getvalue()


def normalise(blob, path):
    data = rezip(blob)
    path.write_bytes(data)
    print(f'{path.name:32} {len(data):>9} bytes')


def layout(prs, name):
    return next(l for l in prs.slide_layouts if l.name == name)


def title_only(prs, text):
    s = prs.slides.add_slide(layout(prs, 'Title Only'))
    s.shapes.title.text = text
    return s


def textbox(slide, x, y, w, h, text, size=18, font=None, bold=False, color=None, align=None):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    if font:
        r.font.name = font
    if color:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    return tb


def bullets(body, items):
    """items: list of (level, text) or (level, [(text, dict(bold/italic/color))])."""
    tf = body.text_frame
    tf.clear()
    first = True
    for level, content in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        runs = content if isinstance(content, list) else [(content, {})]
        for text, fmt in runs:
            r = p.add_run()
            r.text = text
            if fmt.get('bold'):
                r.font.bold = True
            if fmt.get('italic'):
                r.font.italic = True
            if fmt.get('underline'):
                r.font.underline = True
            if fmt.get('color'):
                r.font.color.rgb = fmt['color']


def numbered(paragraph, scheme='arabicPeriod'):
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    au = etree.SubElement(pPr, qn('a:buAutoNum'))
    au.set('type', scheme)


def png_bytes(img, fmt='PNG'):
    b = io.BytesIO()
    img.save(b, fmt, quality=90) if fmt == 'JPEG' else img.save(b, fmt)
    b.seek(0)
    return b


def picture_art(w, h, hue):
    img = Image.new('RGB', (w, h))
    px = img.load()
    for y in range(h):
        for x in range(w):
            px[x, y] = ((hue + x * 180 // w) % 256, 60 + y * 150 // h, 200 - (x + y) * 120 // (w + h))
    d = ImageDraw.Draw(img)
    for i in range(6):
        cx, cy = w * (i + 1) // 7, h // 2 + int(math.sin(i) * h / 4)
        d.ellipse((cx - h // 8, cy - h // 8, cx + h // 8, cy + h // 8), fill=(255, 255, 255), outline=(20, 20, 20), width=3)
    d.rectangle((4, 4, w - 5, h - 5), outline=(0, 0, 0), width=6)
    # A labelled grid so crops are easy to see.
    for gx in range(0, w, w // 4):
        d.line((gx, 0, gx, h), fill=(0, 0, 0), width=2)
    for gy in range(0, h, h // 4):
        d.line((0, gy, w, gy), fill=(0, 0, 0), width=2)
    for i in range(4):
        for j in range(4):
            d.text((i * w // 4 + 10, j * h // 4 + 10), f'{"ABCD"[i]}{j + 1}', fill=(0, 0, 0))
    return img


def transparent_art(size):
    img = Image.new('RGBA', (size, size), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.ellipse((10, 10, size - 10, size - 10), fill=(59, 63, 181, 200))
    d.polygon([(size // 2, 30), (size - 40, size - 40), (40, size - 40)], fill=(232, 163, 23, 255))
    return img


def set_xfrm_attr(shape, **attrs):
    xfrm = shape._element.find('.//' + qn('a:xfrm'))
    for k, v in attrs.items():
        xfrm.set(k, v)


def line_arrow(connector):
    ln = connector.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn('a:tailEnd'))
    tail.set('type', 'triangle')
    tail.set('w', 'med')
    tail.set('len', 'med')


# ---------------------------------------------------------------- 01 title + bullets
def f01():
    prs = new_deck('Title and bullets')
    s = prs.slides.add_slide(layout(prs, 'Title Slide'))
    s.shapes.title.text = 'Quarterly Planning Review'
    s.placeholders[1].text = 'Synthetic corpus deck 01 — title and bullet slides'

    s = prs.slides.add_slide(layout(prs, 'Title and Content'))
    s.shapes.title.text = 'Five levels of bullets'
    bullets(s.placeholders[1], [
        (0, 'Level one: the main point of the slide'),
        (1, 'Level two: a supporting detail'),
        (2, 'Level three: a finer detail'),
        (3, 'Level four: rarely used'),
        (4, 'Level five: the deepest level'),
        (0, [('Mixed runs: ', {}), ('bold', {'bold': True}), (', ', {}), ('italic', {'italic': True}),
             (', ', {}), ('underlined', {'underline': True}), (' and ', {}), ('coloured', {'color': ROSE})]),
        (1, 'Back to level two after a level-one item'),
    ])

    s = prs.slides.add_slide(layout(prs, 'Title and Content'))
    s.shapes.title.text = 'Numbered steps and alignment'
    body = s.placeholders[1]
    bullets(body, [(0, 'Collect the feedback'), (0, 'Group it by theme'), (1, 'Sub-step a'), (1, 'Sub-step b'),
                   (0, 'Decide what to build'), (0, 'Centred paragraph'), (0, 'Right-aligned paragraph')])
    ps = body.text_frame.paragraphs
    for p in (ps[0], ps[1], ps[4]):
        numbered(p)
    for p in (ps[2], ps[3]):
        numbered(p, 'alphaLcParenR')
    ps[5].alignment = PP_ALIGN.CENTER
    ps[6].alignment = PP_ALIGN.RIGHT

    s = prs.slides.add_slide(layout(prs, 'Title and Content'))
    s.shapes.title.text = 'Too much text: PowerPoint shrinks it to fit'
    bullets(s.placeholders[1], [(i % 3, f'Line {i + 1}: a long sentence that keeps going to fill the body placeholder with words') for i in range(16)])

    s = prs.slides.add_slide(layout(prs, 'Two Content'))
    s.shapes.title.text = 'Two columns'
    bullets(s.placeholders[1], [(0, 'Left column'), (1, 'Point A'), (1, 'Point B')])
    bullets(s.placeholders[2], [(0, 'Right column'), (1, 'Point C'), (1, 'Point D')])
    save(prs, '01-title-bullets.pptx')


# ---------------------------------------------------------------- 02 tables
def f02():
    prs = new_deck('Tables')
    s = title_only(prs, 'Table with merged cells and fills')
    rows, cols = 6, 5
    t = s.shapes.add_table(rows, cols, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.2)).table
    data = [['Region', 'Q1', 'Q2', 'Q3', 'Q4'], ['North', '120', '135', '150', '170'], ['South', '90', '95', '110', '118'],
            ['East', '60', '72', '80', '88'], ['West', '140', '150', '149', '160'], ['Total', '410', '452', '489', '536']]
    for r in range(rows):
        for c in range(cols):
            t.cell(r, c).text = data[r][c]
            if c > 0:
                t.cell(r, c).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    t.cell(1, 0).merge(t.cell(2, 0))          # vertical merge
    t.cell(1, 0).text = 'North + South'
    t.cell(5, 1).merge(t.cell(5, 2))          # horizontal merge
    t.cell(5, 1).text = 'H1: 862'
    for c, color in ((3, AMBER), (4, TEAL)):
        cell = t.cell(3, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = color
    t.cell(4, 4).fill.solid()
    t.cell(4, 4).fill.fore_color.rgb = ROSE
    t.cell(4, 4).text_frame.paragraphs[0].runs[0].font.bold = True

    s = title_only(prs, 'Plain table, no style flags, custom widths')
    gf = s.shapes.add_table(4, 3, Inches(1.5), Inches(1.8), Inches(10), Inches(3))
    t = gf.table
    t.first_row = False
    t.horz_banding = False
    t.columns[0].width = Inches(2)
    t.columns[1].width = Inches(5)
    t.columns[2].width = Inches(3)
    for r in range(4):
        for c in range(3):
            cell = t.cell(r, c)
            cell.text = f'Row {r + 1}, column {c + 1}' + (' — a longer cell that wraps onto a second line' if (r, c) == (2, 1) else '')
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xF4, 0xF5, 0xF7) if r % 2 else RGBColor(0xFF, 0xFF, 0xFF)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    save(prs, '02-tables.pptx')


# ---------------------------------------------------------------- 03 images
def f03():
    prs = new_deck('Images')
    art = picture_art(800, 500, 30)
    s = title_only(prs, 'Picture at native aspect, and stretched')
    s.shapes.add_picture(png_bytes(art), Inches(0.6), Inches(1.6), width=Inches(5.8))
    s.shapes.add_picture(png_bytes(art), Inches(6.9), Inches(1.6), width=Inches(5.8), height=Inches(2.2))
    s.shapes.add_picture(png_bytes(transparent_art(300)), Inches(9.5), Inches(4.1), width=Inches(2.8))

    s = title_only(prs, 'Cropped pictures')
    p = s.shapes.add_picture(png_bytes(art), Inches(0.6), Inches(1.6), width=Inches(4))
    p.crop_left, p.crop_right = 0.25, 0.25     # middle half (columns B-C)
    p = s.shapes.add_picture(png_bytes(art), Inches(4.9), Inches(1.6), width=Inches(4))
    p.crop_top, p.crop_bottom = 0.5, 0.0       # bottom half (rows 3-4)
    p = s.shapes.add_picture(png_bytes(art.convert('RGB'), 'JPEG'), Inches(9.2), Inches(1.6), width=Inches(3.5))
    p.crop_left, p.crop_top, p.crop_right, p.crop_bottom = 0.5, 0.25, 0.0, 0.25  # JPEG, right half, middle rows
    textbox(s, Inches(0.6), Inches(6.2), Inches(12), Inches(0.6), 'Left: columns B–C. Middle: rows 3–4. Right (JPEG): columns C–D, rows 2–3.', 14, color=GREY)

    s = prs.slides.add_slide(layout(prs, 'Blank'))
    p = s.shapes.add_picture(png_bytes(art), 0, 0, width=Emu(W), height=Emu(H))  # full-bleed picture
    textbox(s, Inches(0.8), Inches(3), Inches(11.7), Inches(1.4), 'Text over a full-bleed picture', 44, bold=True,
            color=RGBColor(0xFF, 0xFF, 0xFF), align=PP_ALIGN.CENTER)
    save(prs, '03-images.pptx')


# ---------------------------------------------------------------- 04 shapes + groups
def f04():
    prs = new_deck('Shapes and groups')
    s = title_only(prs, 'Preset shapes, fills and lines')
    specs = [
        (MSO_SHAPE.RECTANGLE, 'Rectangle', INDIGO), (MSO_SHAPE.ROUNDED_RECTANGLE, 'Rounded', TEAL),
        (MSO_SHAPE.OVAL, 'Oval', AMBER), (MSO_SHAPE.ISOSCELES_TRIANGLE, 'Triangle', ROSE),
        (MSO_SHAPE.RIGHT_ARROW, 'Arrow', INDIGO), (MSO_SHAPE.STAR_5_POINT, 'Star', AMBER),
        (MSO_SHAPE.CHEVRON, 'Chevron', TEAL), (MSO_SHAPE.HEXAGON, 'Hexagon', ROSE),
        (MSO_SHAPE.CAN, 'Can', GREY), (MSO_SHAPE.CLOUD, 'Cloud', TEAL),
    ]
    for i, (kind, label, color) in enumerate(specs):
        x = Inches(0.6 + (i % 5) * 2.5)
        y = Inches(1.7 + (i // 5) * 2.6)
        sh = s.shapes.add_shape(kind, x, y, Inches(2.1), Inches(1.8))
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
        sh.line.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
        sh.line.width = Pt(2 + (i % 3))
        if i % 4 == 1:
            sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        sh.text_frame.text = label
        sh.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
        sh.text_frame.paragraphs[0].runs[0].font.bold = True

    s = title_only(prs, 'Gradients, rotation, flips, no fill, connectors')
    g = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.7), Inches(3.5), Inches(2))
    g.fill.gradient()
    g.fill.gradient_angle = 45
    g.fill.gradient_stops[0].color.rgb = INDIGO
    g.fill.gradient_stops[1].color.rgb = TEAL
    g.text_frame.text = 'Linear gradient 45°'
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5), Inches(1.9), Inches(3), Inches(1.6))
    r.rotation = 30
    r.fill.solid()
    r.fill.fore_color.rgb = AMBER
    r.text_frame.text = 'Rotated 30°'
    f = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(9), Inches(1.7), Inches(3), Inches(2))
    f.fill.solid()
    f.fill.fore_color.rgb = ROSE
    set_xfrm_attr(f, flipH='1')
    n = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(4.4), Inches(3.5), Inches(2))
    n.fill.background()
    n.line.color.rgb = INDIGO
    n.line.width = Pt(4)
    n.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT
    n.text_frame.text = 'No fill, dotted outline'
    n.text_frame.paragraphs[0].runs[0].font.color.rgb = INDIGO
    a = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.2), Inches(4.6), Inches(1.6), Inches(1.6))
    b = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.6), Inches(4.6), Inches(1.6), Inches(1.6))
    for o in (a, b):
        o.fill.solid()
        o.fill.fore_color.rgb = TEAL
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(6.8), Inches(5.4), Inches(9.6), Inches(5.4))
    c.line.color.rgb = RGBColor(0, 0, 0)
    c.line.width = Pt(3)
    line_arrow(c)
    c2 = s.shapes.add_connector(MSO_CONNECTOR.ELBOW, Inches(6), Inches(4.6), Inches(10.4), Inches(4.0))
    c2.line.color.rgb = GREY
    c2.line.width = Pt(2)

    s = title_only(prs, 'Grouped shapes, nested groups')
    grp = s.shapes.add_group_shape()
    for i, color in enumerate((INDIGO, TEAL, AMBER)):
        sh = grp.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 1.9), Inches(2), Inches(1.7), Inches(1.2))
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
        sh.text_frame.text = f'Group A{i + 1}'
    inner = grp.shapes.add_group_shape()
    for i in range(2):
        sh = inner.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.4 + i * 2.6), Inches(3.6), Inches(1.4), Inches(1.4))
        sh.fill.solid()
        sh.fill.fore_color.rgb = ROSE
        sh.text_frame.text = f'Nested {i + 1}'
    txt = grp.shapes.add_textbox(Inches(0.8), Inches(5.3), Inches(5.5), Inches(0.6))
    txt.text_frame.text = 'A text box inside the group'
    # A second group that is rotated as a whole.
    g2 = s.shapes.add_group_shape()
    for i, color in enumerate((TEAL, INDIGO)):
        sh = g2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8 + i * 1.8), Inches(2.4), Inches(1.6), Inches(2.4))
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
        sh.text_frame.text = f'B{i + 1}'
    g2.rotation = 15
    save(prs, '04-shapes-groups.pptx')


# ---------------------------------------------------------------- 05 charts
def f05():
    prs = new_deck('Charts')
    cd = CategoryChartData()
    cd.categories = ['Q1', 'Q2', 'Q3', 'Q4']
    cd.add_series('2025', (19.2, 21.4, 16.7, 24.1))
    cd.add_series('2026', (22.3, 28.6, 15.2, 30.4))
    s = title_only(prs, 'Clustered column chart')
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.5), cd).chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.plots[0].has_data_labels = True
    ch.plots[0].data_labels.number_format = '0.0'
    ch.plots[0].data_labels.number_format_is_linked = False

    s = title_only(prs, 'Horizontal bar and line charts')
    ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(1.5), Inches(6), Inches(5.3), cd).chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch2 = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(6.8), Inches(1.5), Inches(6), Inches(5.3), cd).chart
    ch2.has_legend = True
    ch2.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch2.legend.include_in_layout = False

    pie = CategoryChartData()
    pie.categories = ['Product', 'Services', 'Licences', 'Other']
    pie.add_series('Revenue mix', (0.46, 0.28, 0.18, 0.08))
    s = title_only(prs, 'Pie and doughnut charts')
    ch = s.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(0.5), Inches(1.5), Inches(6), Inches(5.3), pie).chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.RIGHT
    ch.legend.include_in_layout = False
    ch.plots[0].has_data_labels = True
    dl = ch.plots[0].data_labels
    dl.number_format = '0%'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    ch2 = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(6.8), Inches(1.5), Inches(6), Inches(5.3), pie).chart
    ch2.has_legend = True
    ch2.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch2.legend.include_in_layout = False
    save(prs, '05-charts.pptx')


# ---------------------------------------------------------------- 06 fonts
def set_theme_fonts(prs, major, minor):
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    root.find('.//a:majorFont/a:latin', ns).set('typeface', major)
    root.find('.//a:minorFont/a:latin', ns).set('typeface', minor)
    part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def f06():
    prs = new_deck('Fonts')
    set_theme_fonts(prs, 'Aptos Display', 'Aptos')
    s = title_only(prs, 'Theme fonts: Aptos Display title, Aptos body')
    sample = 'The quick brown fox jumps over the lazy dog 0123456789'
    y = Inches(1.5)
    rows = [(None, 'theme minor (+mn-lt → Aptos)'), ('+mj-lt', 'theme major (+mj-lt → Aptos Display)'),
            ('Calibri', 'Calibri'), ('Calibri Light', 'Calibri Light'), ('Aptos', 'Aptos'), ('Arial', 'Arial'),
            ('Helvetica Neue', 'Helvetica Neue'), ('Times New Roman', 'Times New Roman'), ('Georgia', 'Georgia'),
            ('Courier New', 'Courier New'), ('Ritemark Nowhere Sans', 'Ritemark Nowhere Sans (installed nowhere)')]
    for font, label in rows:
        textbox(s, Inches(0.6), y, Inches(3.8), Inches(0.45), label, 12, color=GREY)
        textbox(s, Inches(4.4), y, Inches(8.5), Inches(0.45), sample, 20, font=font)
        y += Inches(0.52)

    s = title_only(prs, 'Wrapping: the same paragraph in four fonts')
    para = ('Ritemark opens presentations read-only so you can check a deck without leaving your notes. '
            'Line breaks depend on the font metrics, so a missing font moves every line that follows.')
    for i, font in enumerate((None, 'Calibri', 'Arial', 'Ritemark Nowhere Sans')):
        x = Inches(0.5 + (i % 2) * 6.3)
        yy = Inches(1.6 + (i // 2) * 2.7)
        textbox(s, x, yy, Inches(6), Inches(0.4), font or 'theme (Aptos)', 12, color=GREY)
        tb = textbox(s, x, yy + Inches(0.4), Inches(6), Inches(2), para, 18, font=font)
        tb.line.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    save(prs, '06-fonts.pptx')


# ---------------------------------------------------------------- 07 notes
def f07():
    prs = new_deck('Speaker notes')
    for i in range(3):
        s = prs.slides.add_slide(layout(prs, 'Title and Content'))
        s.shapes.title.text = f'Slide {i + 1} has speaker notes'
        bullets(s.placeholders[1], [(0, 'What the audience sees'), (1, 'The notes hold what the speaker says')])
        s.notes_slide.notes_text_frame.text = (f'NOTES {i + 1}: remind the audience of the goal, then walk through '
                                               'the two points. Pause for questions at the end.')
    save(prs, '07-notes.pptx')


# ---------------------------------------------------------------- 08 long deck
def f08():
    prs = new_deck('Forty slides')
    art = picture_art(640, 400, 90)
    cd = CategoryChartData()
    cd.categories = ['A', 'B', 'C', 'D', 'E']
    cd.add_series('Series 1', (4, 7, 3, 8, 5))
    for i in range(40):
        kind = i % 5
        if i == 0:
            s = prs.slides.add_slide(layout(prs, 'Title Slide'))
            s.shapes.title.text = 'A forty-slide deck'
            s.placeholders[1].text = 'For render cost'
            continue
        if kind == 0:
            s = title_only(prs, f'{i + 1}. Chart')
            s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(11), Inches(5.3), cd)
        elif kind == 1:
            s = prs.slides.add_slide(layout(prs, 'Title and Content'))
            s.shapes.title.text = f'{i + 1}. Bullets'
            bullets(s.placeholders[1], [(j % 3, f'Point {j + 1} on slide {i + 1}') for j in range(7)])
        elif kind == 2:
            s = title_only(prs, f'{i + 1}. Table')
            t = s.shapes.add_table(6, 4, Inches(1), Inches(1.6), Inches(11), Inches(4)).table
            for r in range(6):
                for c in range(4):
                    t.cell(r, c).text = f'R{r + 1}C{c + 1}'
        elif kind == 3:
            s = title_only(prs, f'{i + 1}. Shapes')
            for j in range(8):
                sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6 + (j % 4) * 3.1), Inches(1.7 + (j // 4) * 2.5), Inches(2.8), Inches(2.1))
                sh.fill.solid()
                sh.fill.fore_color.rgb = (INDIGO, TEAL, AMBER, ROSE)[j % 4]
                sh.text_frame.text = f'Box {j + 1}'
        else:
            s = title_only(prs, f'{i + 1}. Picture')
            s.shapes.add_picture(png_bytes(art), Inches(2.5), Inches(1.5), width=Inches(8.3))
    save(prs, '08-long-40-slides.pptx')


# ---------------------------------------------------------------- 09 layouts + backgrounds
def f09():
    prs = new_deck('Layouts and backgrounds')
    s = prs.slides.add_slide(layout(prs, 'Section Header'))
    s.shapes.title.text = 'Section header layout'
    s.placeholders[1].text = 'Inherited placeholder positions from the layout'
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(0xE8, 0xEA, 0xF6)

    s = prs.slides.add_slide(layout(prs, 'Comparison'))
    s.shapes.title.text = 'Comparison layout on a gradient background'
    for idx, text in zip([1, 2, 3, 4], ['Option A', 'Faster to ship; fewer features', 'Option B', 'Slower; complete']):
        s.placeholders[idx].text = text
    bg = s.background.fill
    bg.gradient()
    bg.gradient_angle = 90
    bg.gradient_stops[0].color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    bg.gradient_stops[1].color.rgb = RGBColor(0xC9, 0xE9, 0xE4)

    s = prs.slides.add_slide(layout(prs, 'Title and Content'))
    s.shapes.title.text = 'Dark background, light text, a hyperlink'
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(0x1E, 0x21, 0x3A)
    s.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    bullets(s.placeholders[1], [(0, 'White body text on a dark slide'), (0, 'See the product site')])
    for p in s.placeholders[1].text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r = s.placeholders[1].text_frame.paragraphs[1].runs[0]
    r.hyperlink.address = 'https://example.com/'

    s = prs.slides.add_slide(layout(prs, 'Content with Caption'))
    s.shapes.title.text = 'Content with caption layout'
    for ph in s.placeholders:
        if ph.placeholder_format.idx == 1:
            ph.text = 'Main content area'
        elif ph.placeholder_format.idx == 2:
            ph.text = 'A caption that explains the content next to it.'
    save(prs, '09-layouts-backgrounds.pptx')


# ---------------------------------------------------------------- failure fixtures
def write_bomb(src_path, part, anchor, out_path, target=512 * 1024 * 1024):
    """Copy an Office package, padding `part` with whitespace inside its root element until it
    inflates to `target` bytes. Structurally valid; about 0.5 MB on disk."""
    src = zipfile.ZipFile(src_path)
    head, sep, tail = src.read(part).partition(anchor)
    assert sep, f'{anchor!r} not in {part}'
    chunk = b' ' * (8 * 1024 * 1024)
    with zipfile.ZipFile(out_path, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as dst:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            if info.filename == part:
                with dst.open(zi, 'w', force_zip64=False) as w:
                    w.write(head)
                    written = len(head)
                    while written + len(chunk) <= target - len(sep) - len(tail):
                        w.write(chunk)
                        written += len(chunk)
                    w.write(sep + tail)
            else:
                dst.writestr(zi, src.read(info.filename))
    print(f'{out_path.name:32} {out_path.stat().st_size:>9} bytes')


def write_lying_bomb(bomb_path, part, out_path, lie=4096):
    """The same bomb, but the archive declares `part` as `lie` bytes: in the local header, the
    central directory and, if present, the data descriptor. A check that trusts declared sizes
    lets it through."""
    data = bytearray(bomb_path.read_bytes())
    info = zipfile.ZipFile(io.BytesIO(bytes(data))).getinfo(part)
    name = part.encode()
    # local header: uncompressed size at offset 22; central directory entry: offset 24
    lh = info.header_offset
    assert data[lh:lh + 4] == b'PK\x03\x04'
    flags = int.from_bytes(data[lh + 6:lh + 8], 'little')
    data[lh + 22:lh + 26] = lie.to_bytes(4, 'little')
    pos = 0
    while True:
        pos = data.find(b'PK\x01\x02', pos)
        if pos < 0:
            raise SystemExit('central directory entry not found')
        nlen = int.from_bytes(data[pos + 28:pos + 30], 'little')
        if bytes(data[pos + 46:pos + 46 + nlen]) == name:
            data[pos + 24:pos + 28] = lie.to_bytes(4, 'little')
            break
        pos += 4
    # a data descriptor (flag bit 3) repeats the size after the data; zipfile writes one when streaming
    if flags & 0x08:
        dd = data.find(b'PK\x07\x08', lh + 30 + len(name) + info.compress_size - 16)
        if dd > 0:
            data[dd + 12:dd + 16] = lie.to_bytes(4, 'little')
    out_path.write_bytes(bytes(data))
    print(f'{out_path.name:32} {len(data):>9} bytes')


def failures():
    (OUT / 'f1-not-a-zip.pptx').write_text('This is a plain text file with a .pptx name.\n' * 20)
    whole = (OUT / '01-title-bullets.pptx').read_bytes()
    (OUT / 'f2-truncated.pptx').write_bytes(whole[: len(whole) // 2])

    # f3: a structurally valid deck whose slide1.xml inflates to 512 MiB.
    bomb = OUT / 'f3-decompression-bomb.pptx'
    write_bomb(OUT / '01-title-bullets.pptx', 'ppt/slides/slide1.xml', b'<p:cSld', bomb)
    # f4: the same bomb, declaring slide1.xml as 4 KiB.
    write_lying_bomb(bomb, 'ppt/slides/slide1.xml', OUT / 'f4-bomb-lying-size.pptx')
    # f5 (password-protected) is written by PowerPoint itself: ppt-password.sh.

    # f6: the lying bomb as a Word file, from Sprint 124's Word-saved fixture 01. Sprint 124's
    # host check reads declared sizes only; this shows whether a lying .docx gets past it.
    word = HERE.parents[2] / 'sprint-124-word-preview-fidelity/research/corpus/word/01-headings-styles.docx'
    scratch = HERE / 'scratch'
    scratch.mkdir(exist_ok=True)
    docx_bomb = scratch / 'f6-source-bomb.docx'
    write_bomb(word, 'word/document.xml', b'<w:body', docx_bomb)
    write_lying_bomb(docx_bomb, 'word/document.xml', OUT / 'f6-docx-bomb-lying-size.docx')
    docx_bomb.unlink()


if __name__ == '__main__':
    for f in (f01, f02, f03, f04, f05, f06, f07, f08, f09):
        f()
    failures()
